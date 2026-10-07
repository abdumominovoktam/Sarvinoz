import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from django.db.models import Avg, Max, Count, Q
from django.shortcuts import get_object_or_404
from rest_framework import status, generics
from rest_framework.views import APIView
from rest_framework.response import Response

from apps.users.models import User, Group
from apps.users.serializers import UserSerializer, GroupSerializer
from apps.users.permissions import IsAdminRole
from apps.games.models import Game, Question
from apps.games.serializers import AdminQuestionSerializer, GameListSerializer
from apps.results.models import GameResult, ViolationLog
from apps.results.serializers import GameResultSerializer, ViolationLogSerializer
from apps.leaderboard.models import Leaderboard
from apps.leaderboard.services import recalculate_all_rankings


class AdminDashboardStatsView(APIView):
    permission_classes = [IsAdminRole]

    def get(self, request):
        recalculate_all_rankings()

        students_qs = User.objects.filter(role='student')
        total_students = students_qs.count()
        active_students = students_qs.filter(is_blocked=False).count()

        lb_qs = Leaderboard.objects.filter(user__role='student').select_related('user', 'group')
        completed_students = lb_qs.filter(completed_games__gte=8).count()

        agg = lb_qs.aggregate(
            avg_score=Avg('total_score'),
            max_score=Max('total_score'),
            avg_time=Avg('total_time_spent'),
            avg_acc=Avg('accuracy'),
        )

        fastest_entry = (
            lb_qs.filter(completed_games__gte=8, total_time_spent__gt=0)
            .order_by('total_time_spent', '-total_score')
            .first()
        )
        if not fastest_entry:
            fastest_entry = lb_qs.filter(total_time_spent__gt=0).order_by('total_time_spent').first()

        top_scorer_entry = lb_qs.order_by('-total_score', 'total_time_spent').first()

        most_violations_user = (
            User.objects.filter(role='student')
            .annotate(v_count=Count('violations'))
            .filter(v_count__gt=0)
            .order_by('-v_count')
            .first()
        )

        # Score distribution buckets
        buckets = [
            {'range': '0 - 300', 'min': 0, 'max': 300, 'count': 0},
            {'range': '301 - 500', 'min': 301, 'max': 500, 'count': 0},
            {'range': '501 - 650', 'min': 501, 'max': 650, 'count': 0},
            {'range': '651 - 750', 'min': 651, 'max': 750, 'count': 0},
            {'range': '751 - 867', 'min': 751, 'max': 9999, 'count': 0},
        ]
        for score_val in lb_qs.values_list('total_score', flat=True):
            for b in buckets:
                if b['min'] <= score_val <= b['max']:
                    b['count'] += 1
                    break

        # Per-game analytics
        games = Game.objects.all().order_by('order')
        game_analytics = []
        for g in games:
            g_res = GameResult.objects.filter(game=g)
            g_agg = g_res.aggregate(
                completions=Count('id'),
                avg_score=Avg('score'),
                avg_time=Avg('time_spent'),
                avg_acc=Avg('accuracy'),
            )
            game_analytics.append({
                'id': g.id,
                'order': g.order,
                'title': g.title,
                'max_score': g.max_score,
                'duration_seconds': g.duration_seconds,
                'completions': g_agg['completions'] or 0,
                'completion_rate': round(((g_agg['completions'] or 0) / max(1, total_students)) * 100.0, 1),
                'average_score': round(g_agg['avg_score'] or 0.0, 1),
                'average_time': round(g_agg['avg_time'] or 0.0, 1),
                'average_accuracy': round(g_agg['avg_acc'] or 0.0, 1),
            })

        # Group comparison
        group_comparison = []
        for grp in Group.objects.all():
            g_lb = lb_qs.filter(group=grp)
            cnt = g_lb.count()
            if cnt == 0:
                continue
            g_a = g_lb.aggregate(
                avg_score=Avg('total_score'),
                avg_time=Avg('total_time_spent'),
                avg_acc=Avg('accuracy'),
            )
            avg_s = round(g_a['avg_score'] or 0.0, 1)
            group_comparison.append({
                'group_name': grp.name,
                'student_count': cnt,
                'average_score': avg_s,
                'percentage': round((avg_s / 867.0) * 100.0, 1),
                'average_time': round(g_a['avg_time'] or 0.0, 1),
                'average_accuracy': round(g_a['avg_acc'] or 0.0, 1),
            })
        group_comparison.sort(key=lambda x: -x['average_score'])

        return Response({
            'kpis': {
                'total_students': total_students,
                'active_students': active_students,
                'completed_students': completed_students,
                'average_score': round(agg['avg_score'] or 0.0, 1),
                'highest_score': agg['max_score'] or 0,
                'top_scorer_name': top_scorer_entry.user.full_name if top_scorer_entry else '—',
                'average_time_seconds': round(agg['avg_time'] or 0.0, 1),
                'average_accuracy': round(agg['avg_acc'] or 0.0, 1),
                'fastest_finisher': {
                    'name': fastest_entry.user.full_name if fastest_entry else '—',
                    'group': fastest_entry.group.name if (fastest_entry and fastest_entry.group) else '—',
                    'time_spent': fastest_entry.total_time_spent if fastest_entry else 0,
                    'score': fastest_entry.total_score if fastest_entry else 0,
                },
                'most_violations': {
                    'name': most_violations_user.full_name if most_violations_user else 'Yo‘q',
                    'group': most_violations_user.group_name if most_violations_user else '—',
                    'count': most_violations_user.v_count if most_violations_user else 0,
                },
                'total_violations_count': ViolationLog.objects.count(),
            },
            'score_distribution': buckets,
            'game_analytics': game_analytics,
            'group_comparison': group_comparison,
        })


class AdminStudentListView(APIView):
    permission_classes = [IsAdminRole]

    def get(self, request):
        search = request.query_params.get('search', '').strip()
        group_name = request.query_params.get('group', '').strip()

        qs = User.objects.filter(role='student').select_related('group', 'profile').order_by('-created_at')
        if search:
            qs = qs.filter(
                Q(first_name__icontains=search)
                | Q(last_name__icontains=search)
                | Q(email__icontains=search)
            )
        if group_name and group_name != 'all':
            qs = qs.filter(group__name__iexact=group_name)

        lb_map = {lb.user_id: lb for lb in Leaderboard.objects.all()}
        v_counts = dict(
            User.objects.filter(role='student')
            .annotate(vc=Count('violations'))
            .values_list('id', 'vc')
        )

        data = []
        for u in qs:
            lb = lb_map.get(u.id)
            data.append({
                **UserSerializer(u).data,
                'total_score': lb.total_score if lb else 0,
                'completed_games': lb.completed_games if lb else 0,
                'total_time_spent': lb.total_time_spent if lb else 0,
                'accuracy': lb.accuracy if lb else 0.0,
                'rank': lb.rank if lb else None,
                'violations_count': v_counts.get(u.id, 0),
            })

        return Response({'students': data})


class AdminToggleBlockStudentView(APIView):
    permission_classes = [IsAdminRole]

    def post(self, request, user_id: int):
        student = get_object_or_404(User, id=user_id, role='student')
        student.is_blocked = not student.is_blocked
        student.save(update_fields=['is_blocked'])
        return Response({
            'message': f"{student.full_name} {'bloklandi' if student.is_blocked else 'blokdan chiqarildi'}.",
            'is_blocked': student.is_blocked,
            'user_id': student.id,
        })


class AdminGroupListCreateView(generics.ListCreateAPIView):
    permission_classes = [IsAdminRole]
    queryset = Group.objects.all().order_by('name')
    serializer_class = GroupSerializer


class AdminGroupDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAdminRole]
    queryset = Group.objects.all()
    serializer_class = GroupSerializer


class AdminGameListView(generics.ListAPIView):
    permission_classes = [IsAdminRole]
    queryset = Game.objects.all().order_by('order')
    serializer_class = GameListSerializer


class AdminGameUpdateView(APIView):
    permission_classes = [IsAdminRole]

    def patch(self, request, game_id: int):
        game = get_object_or_404(Game, id=game_id)
        for field in ['title', 'description', 'duration_seconds', 'max_score', 'is_active']:
            if field in request.data:
                setattr(game, field, request.data[field])
        game.save()
        return Response(GameListSerializer(game).data)

    def put(self, request, game_id: int):
        return self.patch(request, game_id)


class AdminQuestionListCreateView(generics.ListCreateAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = AdminQuestionSerializer

    def get_queryset(self):
        qs = Question.objects.select_related('game').order_by('game__order', '-id')
        game_id = self.request.query_params.get('game')
        difficulty = self.request.query_params.get('difficulty')
        search = self.request.query_params.get('search', '').strip()
        if game_id:
            qs = qs.filter(game_id=game_id)
        if difficulty and difficulty != 'all':
            qs = qs.filter(difficulty__iexact=difficulty)
        if search:
            qs = qs.filter(question_text__icontains=search)
        return qs


class AdminQuestionDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAdminRole]
    queryset = Question.objects.select_related('game').all()
    serializer_class = AdminQuestionSerializer


class AdminResultsAndViolationsView(APIView):
    permission_classes = [IsAdminRole]

    def get(self, request):
        recent_results = GameResult.objects.select_related('user', 'user__group', 'game').order_by('-completed_at')[:100]
        violations = ViolationLog.objects.select_related('user', 'user__group', 'game').order_by('-timestamp')[:100]
        return Response({
            'game_results': GameResultSerializer(recent_results, many=True).data,
            'violations': ViolationLogSerializer(violations, many=True).data,
        })


class AdminExportResultsView(APIView):
    permission_classes = [IsAdminRole]

    def get(self, request):
        recalculate_all_rankings()
        export_format = request.query_params.get('export_format', 'csv').lower()
        entries = list(
            Leaderboard.objects.select_related('user', 'group')
            .filter(user__role='student')
            .order_by('rank')
        )

        headers = [
            'O‘rin (Rank)',
            'Familiya',
            'Ism',
            'Guruh',
            'Email',
            'Umumiy Ball (max 867)',
            'Bajarilgan O‘yinlar (8 dan)',
            'Sarflangan Vaqt (soniya)',
            'Sarflangan Vaqt (MM:SS)',
            'O‘rtacha Vaqt (soniya)',
            'Aniqlik (%)',
            'Guruhdagi O‘rni',
        ]

        rows = []
        for e in entries:
            mins = e.total_time_spent // 60
            secs = e.total_time_spent % 60
            rows.append([
                e.rank,
                e.user.last_name,
                e.user.first_name,
                e.group.name if e.group else e.user.group_name,
                e.user.email,
                e.total_score,
                f"{e.completed_games}/8",
                e.total_time_spent,
                f"{mins:02d}:{secs:02d}",
                round(e.average_time, 1),
                f"{e.accuracy}%",
                e.group_rank,
            ])

        if export_format in ('excel', 'xlsx'):
            wb = Workbook()
            ws = wb.active
            ws.title = "Talabalar Reytingi"

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="007A63", end_color="007A63", fill_type="solid")

            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row in rows:
                ws.append(row)

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(14, max_len + 4)

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)

            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            response['Content-Disposition'] = 'attachment; filename="sanoq_sistemalari_reyting.xlsx"'
            return response

        # Default CSV export (with UTF-8 BOM so Excel opens Uzbek characters cleanly)
        response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
        response['Content-Disposition'] = 'attachment; filename="sanoq_sistemalari_reyting.csv"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
        return response
